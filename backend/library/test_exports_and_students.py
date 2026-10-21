"""Catalogue / loan exports and the bulk student import."""
import csv
import io
import re
import tempfile
from datetime import datetime, timezone as dt_timezone
from pathlib import Path

from django.contrib.auth.models import User
from django.core import mail
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from . import services
from .exports import BOOK_COLUMNS, LOAN_COLUMNS, safe_cell
from .importing import read_rows
from .models import ActivityLog, Book, BorrowRecord, Category, Role, Status, UserProfile
from .tests import make_book, make_user

BOM = b"\xef\xbb\xbf"


def csv_rows(res) -> list[list[str]]:
    body = res.content if not res.streaming else b"".join(res.streaming_content)
    assert body.startswith(BOM), "CSV must start with a UTF-8 BOM"
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


def xlsx_rows(res) -> list[list]:
    sheet = load_workbook(io.BytesIO(res.content), read_only=True).worksheets[0]
    return [list(r) for r in sheet.iter_rows(values_only=True)]


class ExportBase(TestCase):
    def setUp(self):
        self.admin = make_user("a@x.com", Role.ADMIN)
        self.staff = make_user("s@x.com", Role.STAFF)
        self.member = make_user("m@x.com")
        fiction = Category.objects.create(name="Fiction", color="#111111")
        history = Category.objects.create(name="History", color="#222222")
        self.dune = Book.objects.create(
            isbn="9780441172719", title="Dune", author="Frank Herbert", category=fiction, publisher="Ace",
            year=1990, total_copies=3, available_copies=3, shelf_location="FIC-A-09", description="Sci-fi, classic")
        self.sapiens = Book.objects.create(
            isbn="9780062316097", title="Sapiens", author="Harari", category=history, total_copies=2,
            available_copies=2)
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

    def get(self, path, **params):
        return self.client.get(path, params)


class BookExportTests(ExportBase):
    def test_csv_has_the_import_columns_then_the_stock_columns(self):
        res = self.get("/api/exports/books", format="csv")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res["Content-Type"].startswith("text/csv"))
        self.assertRegex(res["Content-Disposition"],
                         r'^attachment; filename="library-books-\d{4}-\d{2}-\d{2}\.csv"$')
        rows = csv_rows(res)
        self.assertEqual(rows[0], ["isbn", "title", "author", "category", "publisher", "year", "total_copies",
                                   "shelf_location", "description", "available_copies", "free_copies"])
        self.assertEqual(rows[0], BOOK_COLUMNS)
        self.assertEqual(rows[1], ["9780441172719", "Dune", "Frank Herbert", "Fiction", "Ace", "1990", "3",
                                   "FIC-A-09", "Sci-fi, classic", "3", "3"])
        self.assertEqual(len(rows), 3)

    def test_xlsx_headers_and_a_row(self):
        res = self.get("/api/exports/books", format="xlsx")
        self.assertEqual(res.status_code, 200)
        self.assertIn("spreadsheetml", res["Content-Type"])
        self.assertRegex(res["Content-Disposition"], r'filename="library-books-\d{4}-\d{2}-\d{2}\.xlsx"')
        rows = xlsx_rows(res)
        self.assertEqual(rows[0], BOOK_COLUMNS)
        self.assertEqual(rows[1][:3], ["9780441172719", "Dune", "Frank Herbert"])
        self.assertEqual(rows[1][5], 1990)
        self.assertEqual(rows[2][5], None)  # no year

    def test_free_copies_leave_out_the_holds(self):
        loan = services.create_request(self.member, self.dune)
        services.approve(self.staff, loan)  # held: still on the shelf, no longer free
        row = csv_rows(self.get("/api/exports/books", format="csv"))[1]
        self.assertEqual(row[-2:], ["3", "2"])

    def test_filters_apply(self):
        rows = csv_rows(self.get("/api/exports/books", format="csv", search="sapi"))
        self.assertEqual([r[1] for r in rows[1:]], ["Sapiens"])
        rows = csv_rows(self.get("/api/exports/books", format="csv", category="Fiction"))
        self.assertEqual([r[1] for r in rows[1:]], ["Dune"])

    def test_format_defaults_to_csv_and_rejects_others(self):
        self.assertEqual(csv_rows(self.get("/api/exports/books"))[0][0], "isbn")
        res = self.get("/api/exports/books", format="pdf")
        self.assertEqual((res.status_code, res.data["code"]), (400, "export_bad_format"))

    def test_formula_injection_is_defused_in_both_formats(self):
        self.dune.title = '=HYPERLINK("x")'
        self.dune.author = "@SUM(1)"
        self.dune.publisher = "+1"
        self.dune.description = "-2+3"
        self.dune.shelf_location = "\t=1"
        self.dune.save()
        for rows in (csv_rows(self.get("/api/exports/books", format="csv")),
                     xlsx_rows(self.get("/api/exports/books", format="xlsx"))):
            dune = next(r for r in rows[1:] if r[0] == "9780441172719")
            self.assertEqual(dune[1], "'=HYPERLINK(\"x\")")
            self.assertEqual((dune[2], dune[4], dune[8], dune[7]), ("'@SUM(1)", "'+1", "'-2+3", "'\t=1"))
            self.assertEqual(dune[3], "Fiction")  # ordinary text is untouched

    def test_safe_cell_only_touches_text_that_looks_like_a_formula(self):
        self.assertEqual(safe_cell("=1"), "'=1")
        self.assertEqual(safe_cell("\r=1"), "'\r=1")
        self.assertEqual(safe_cell("a=1"), "a=1")
        self.assertEqual(safe_cell(-5), -5)
        self.assertEqual(safe_cell(""), "")
        self.assertIsNone(safe_cell(None))

    def test_arabic_survives_with_the_bom(self):
        self.dune.title = "كتاب الرمل"
        self.dune.save()
        res = self.get("/api/exports/books", format="csv")
        rows = csv_rows(res)
        self.assertEqual(next(r for r in rows if r[0] == "9780441172719")[1], "كتاب الرمل")

    def test_round_trip_through_the_import_changes_nothing(self):
        for fmt, name in (("csv", "books.csv"), ("xlsx", "books.xlsx")):
            res = self.get("/api/exports/books", format=fmt)
            content = b"".join(res.streaming_content) if res.streaming else res.content
            rows = read_rows(name, content)
            report = self.client.post("/api/books/import", {"rows": rows, "mode": "update", "dryRun": False},
                                      format="json").data
            self.assertEqual(report["summary"]["error"], 0, report["rows"])
            self.assertEqual(report["summary"]["unchanged"], 2, report["rows"])
            self.assertEqual(report["summary"]["update"], 0, report["rows"])
            # the stock columns are ignored by the import
            self.assertEqual({h for h, f in report["columns"].items() if f is None},
                             {"available_copies", "free_copies"})
        self.dune.refresh_from_db()
        self.assertEqual((self.dune.total_copies, self.dune.available_copies, self.dune.year), (3, 3, 1990))

    def test_members_and_anonymous_users_are_refused(self):
        for who in (self.member, None):
            client = APIClient()
            if who:
                client.force_authenticate(who)
            self.assertEqual(client.get("/api/exports/books?format=csv").status_code, 403 if who else 401)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.get("/api/exports/books", format="csv").status_code, 200)


class LoanExportTests(ExportBase):
    def setUp(self):
        super().setUp()
        self.old = BorrowRecord.objects.create(user=self.member, book=self.dune, status=Status.RETURNED,
                                               renewal_count=1)
        self.mid = BorrowRecord.objects.create(user=self.member, book=self.sapiens, status=Status.BORROWED)
        self.new = BorrowRecord.objects.create(user=self.member, book=self.dune, status=Status.PENDING)
        for record, day in ((self.old, "2026-01-05 10:30"), (self.mid, "2026-02-10 23:30"),
                            (self.new, "2026-03-15 08:00")):
            BorrowRecord.objects.filter(pk=record.pk).update(
                requested_at=datetime.strptime(day, "%Y-%m-%d %H:%M").replace(tzinfo=dt_timezone.utc))
        BorrowRecord.objects.filter(pk=self.mid.pk).update(
            borrowed_at=datetime(2026, 2, 11, 9, 0, tzinfo=dt_timezone.utc),
            due_date=datetime(2026, 2, 25, 9, 0, tzinfo=dt_timezone.utc))

    def titles_and_status(self, res):
        return [(r[3], r[6]) for r in csv_rows(res)[1:]]

    def test_columns_and_a_row(self):
        rows = csv_rows(self.get("/api/exports/loans", format="csv", status="BORROWED"))
        self.assertEqual(rows[0], LOAN_COLUMNS)
        self.assertEqual(rows[1], [
            "m@x.com", "M", "m@x.com", self.sapiens.title, "9780062316097", "",
            "BORROWED", "2026-02-10 23:30", "2026-02-11 09:00", "2026-02-25 09:00", "", "0"])
        self.assertEqual(len(rows), 2)

    def test_filter_by_status_list(self):
        res = self.get("/api/exports/loans", format="csv", status="RETURNED,PENDING")
        self.assertEqual(self.titles_and_status(res), [("Dune", "PENDING"), ("Dune", "RETURNED")])

    def test_filter_by_requested_at_range_inclusive(self):
        res = self.get("/api/exports/loans", format="csv", **{"from": "2026-02-01", "to": "2026-02-10"})
        self.assertEqual(self.titles_and_status(res), [("Sapiens", "BORROWED")])  # 23:30 on the last day counts
        res = self.get("/api/exports/loans", format="csv", **{"from": "2026-02-11"})
        self.assertEqual(self.titles_and_status(res), [("Dune", "PENDING")])
        res = self.get("/api/exports/loans", format="csv", **{"to": "2026-01-05"})
        self.assertEqual(self.titles_and_status(res), [("Dune", "RETURNED")])

    @override_settings(LIBRARY_TIME_ZONE="Africa/Algiers")  # UTC+1 all year
    def test_dates_are_in_the_library_time_zone(self):
        row = csv_rows(self.get("/api/exports/loans", format="csv", status="BORROWED"))[1]
        self.assertEqual((row[7], row[8], row[9]), ("2026-02-11 00:30", "2026-02-11 10:00", "2026-02-25 10:00"))
        # the range is read in the same zone: 23:30 UTC on the 10th is already the 11th in Algiers
        res = self.get("/api/exports/loans", format="csv", **{"from": "2026-02-11", "to": "2026-02-11"})
        self.assertEqual(self.titles_and_status(res), [("Sapiens", "BORROWED")])

    def test_xlsx(self):
        rows = xlsx_rows(self.get("/api/exports/loans", format="xlsx", status="RETURNED"))
        self.assertEqual(rows[0], LOAN_COLUMNS)
        self.assertEqual((rows[1][3], rows[1][11]), ("Dune", 1))

    def test_injection_in_a_member_name(self):
        User.objects.filter(pk=self.member.pk).update(first_name="=cmd|' /C calc'!A0", last_name="X")
        for rows in (csv_rows(self.get("/api/exports/loans", format="csv")),
                     xlsx_rows(self.get("/api/exports/loans", format="xlsx"))):
            self.assertTrue(rows[1][0].startswith("'=cmd"))

    def test_bad_dates_and_permissions(self):
        res = self.get("/api/exports/loans", format="csv", **{"from": "yesterday"})
        self.assertEqual((res.status_code, res.data["code"]), (400, "export_bad_date"))
        client = APIClient()
        client.force_authenticate(self.member)
        self.assertEqual(client.get("/api/exports/loans?format=csv").status_code, 403)


# --------------------------------------------------------------------- student import

PAST = "shelf-mark-HIS-42"


class StudentImportBase(TestCase):
    def setUp(self):
        self.admin = make_user("a@x.com", Role.ADMIN)
        self.staff = make_user("s@x.com", Role.STAFF)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def post(self, rows, **body):
        body.setdefault("dryRun", True)
        return self.client.post("/api/users/import", {"rows": rows, **body}, format="json")


class StudentImportTests(StudentImportBase):
    def test_preview_writes_and_sends_nothing(self):
        users = User.objects.count()
        res = self.post([{"name": "Nadia Salem", "email": "Nadia@Uni.dz", "student_id": "S100"}])
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["rows"][0]["status"], "create")
        self.assertEqual(res.data["rows"][0]["email"], "nadia@uni.dz")
        self.assertEqual(res.data["summary"], {"total": 1, "create": 1, "exists": 0, "error": 0})
        self.assertEqual(User.objects.count(), users)
        self.assertEqual(len(mail.outbox), 0)
        self.assertFalse(ActivityLog.objects.filter(action="STUDENTS_IMPORTED").exists())

    def test_import_creates_unverified_members_and_one_welcome_email_each(self):
        rows = [
            {"name": "Nadia Salem", "email": "nadia@uni.dz", "student_id": "S100", "phone": "0555 12 34", "language": "fr"},
            {"first_name": "Omar", "last_name": "Ben Ali", "email": "omar@uni.dz", "student_id": 101},
            {"Nom": "Haddad", "Prénom": "Lina", "Courriel": "lina@uni.dz", "Matricule": "S102", "Langue": "ar"},
        ]
        res = self.post(rows, dryRun=False, filename="students.csv")
        self.assertEqual(res.data["summary"]["create"], 3, res.data["rows"])
        self.assertEqual(res.data["emailsSent"], 3)

        nadia = User.objects.get(email="nadia@uni.dz")
        self.assertEqual(nadia.username, "nadia@uni.dz")
        self.assertFalse(nadia.has_usable_password())
        self.assertEqual((nadia.first_name, nadia.last_name), ("Nadia", "Salem"))
        p = nadia.profile
        self.assertEqual((p.role, p.email_verified, p.language, p.student_or_staff_id, p.phone_number, p.active),
                         (Role.MEMBER, False, "fr", "S100", "0555 12 34", True))
        omar = User.objects.get(email="omar@uni.dz")
        self.assertEqual((omar.last_name, omar.profile.language, omar.profile.student_or_staff_id), ("Ben Ali", "ar", "101"))
        lina = User.objects.get(email="lina@uni.dz")  # "Nom" next to "Prénom" is the family name
        self.assertEqual((lina.first_name, lina.last_name, lina.profile.language), ("Lina", "Haddad", "ar"))

        self.assertEqual(sorted(m.to[0] for m in mail.outbox), ["lina@uni.dz", "nadia@uni.dz", "omar@uni.dz"])
        by_to = {m.to[0]: m for m in mail.outbox}
        self.assertIn("Votre compte", by_to["nadia@uni.dz"].subject)
        self.assertIn("Mot de passe oublié", by_to["nadia@uni.dz"].body)
        self.assertIn("حسابك", by_to["omar@uni.dz"].subject)
        self.assertIn("نسيت كلمة المرور", by_to["omar@uni.dz"].body)  # no language column: Arabic, the default
        self.assertIn("نسيت كلمة المرور", by_to["lina@uni.dz"].body)
        entry = ActivityLog.objects.get(action="STUDENTS_IMPORTED")
        self.assertEqual((entry.params["created"], entry.params["source"]), (3, "students.csv"))
        self.assertEqual(ActivityLog.objects.filter(action="STUDENTS_IMPORTED").count(), 1)

    def test_each_error_code(self):
        make_user("taken@x.com")  # student id "TAKEN"
        make_user("m@x.com")
        rows = [
            {"name": "Ok One", "email": "one@uni.dz", "student_id": "S1"},                       # 1 create
            {"name": "", "email": "two@uni.dz", "student_id": "S2"},                              # name_missing
            {"name": "Three", "email": "", "student_id": "S3"},                                   # email_missing
            {"name": "Four", "email": "not-an-email", "student_id": "S4"},                        # email_invalid
            {"name": "Five", "email": "five@uni.dz", "student_id": ""},                           # student_id_missing
            {"name": "Six", "email": "ONE@uni.dz", "student_id": "S6"},                           # email_duplicate (row 1)
            {"name": "Seven", "email": "seven@uni.dz", "student_id": "s1"},                       # student_id_duplicate (row 1)
            {"name": "Eight", "email": "eight@uni.dz", "student_id": "TAKEN"},                    # student_id_taken
            {"name": "Nine", "email": "nine@uni.dz", "student_id": "S9", "language": "klingon"},  # language_invalid
            {"name": "Ten", "email": "m@x.com", "student_id": "S10"},                             # exists
            {"name": "Eleven", "email": "e" * 145 + "@uni.dz", "student_id": "S11"},              # email_invalid (too long)
            {"name": "Twelve", "email": "twelve@uni.dz", "student_id": "S" * 31},                 # student_id_too_long
            {"name": "X" * 151, "email": "thirteen@uni.dz", "student_id": "S13"},                 # name_too_long
            {"name": "Fourteen", "email": "fourteen@uni.dz", "student_id": "S14", "phone": "1" * 41},  # phone_invalid
        ]
        rows_by_no = {r["row"]: r for r in self.post(rows).data["rows"]}
        codes = {n: [e["code"] for e in r["errors"]] for n, r in rows_by_no.items()}
        self.assertEqual(codes[2], ["name_missing"])
        self.assertEqual(codes[3], ["email_missing"])
        self.assertEqual(codes[4], ["email_invalid"])
        self.assertEqual(codes[5], ["student_id_missing"])
        self.assertEqual(codes[6], ["email_duplicate"])
        self.assertEqual(rows_by_no[6]["errors"][0]["row"], 1)
        self.assertEqual(codes[7], ["student_id_duplicate"])
        self.assertEqual(rows_by_no[7]["errors"][0]["row"], 1)
        self.assertEqual(codes[8], ["student_id_taken"])
        self.assertEqual(codes[9], ["language_invalid"])
        self.assertEqual((rows_by_no[10]["status"], codes[10]), ("exists", []))
        self.assertEqual(codes[11], ["email_invalid"])
        self.assertEqual(codes[12], ["student_id_too_long"])
        self.assertEqual(codes[13], ["name_too_long"])
        self.assertEqual(codes[14], ["phone_invalid"])
        self.assertEqual(rows_by_no[1]["status"], "create")

    def test_existing_email_is_skipped_not_an_error(self):
        make_user("old@x.com")
        res = self.post([
            {"name": "Old Timer", "email": "OLD@x.com", "student_id": "WHATEVER"},
            {"name": "New One", "email": "new@x.com", "student_id": "N1"},
        ], dryRun=False)
        self.assertEqual([r["status"] for r in res.data["rows"]], ["exists", "create"])
        self.assertEqual(res.data["summary"], {"total": 2, "create": 1, "exists": 1, "error": 0})
        self.assertEqual(res.data["rows"][0]["errors"], [])
        self.assertEqual([m.to[0] for m in mail.outbox], ["new@x.com"])
        self.assertEqual(UserProfile.objects.get(user__email="old@x.com").student_or_staff_id, "OLD")

    def test_error_rows_are_skipped_and_the_rest_is_saved(self):
        res = self.post([
            {"name": "Good", "email": "good@x.com", "student_id": "G1"},
            {"name": "Bad", "email": "bad", "student_id": "B1"},
        ], dryRun=False)
        self.assertEqual((res.data["summary"]["create"], res.data["summary"]["error"]), (1, 1))
        self.assertTrue(User.objects.filter(email="good@x.com").exists())
        self.assertFalse(User.objects.filter(email="bad").exists())

    def test_blank_lines_ignored_missing_columns_and_limits(self):
        res = self.post([{"name": "", "email": "", "student_id": ""}])
        self.assertEqual(res.data["summary"]["total"], 0)
        res = self.post([{"Name": "A", "Mail": "a@x.com"}])
        self.assertEqual(res.data["missingColumns"], ["studentId"])
        self.assertEqual(self.post("nope").data["code"], "import_bad_rows")
        res = self.post([{}] * 5001)
        self.assertEqual((res.status_code, res.data["code"], res.data["params"]), (400, "import_too_many_rows", {"max": 5000}))

    def test_headers_in_french_and_arabic(self):
        res = self.post([{"الاسم": "ليلى بن علي", "البريد الإلكتروني": "layla@uni.dz", "الرقم الجامعي": "A1", "اللغة": "العربية"}])
        row = res.data["rows"][0]
        self.assertEqual((row["status"], row["name"], row["language"]), ("create", "ليلى بن علي", "ar"))

    def test_only_admins(self):
        for user, expected in ((self.staff, 403), (make_user("m2@x.com"), 403), (None, 401)):
            client = APIClient()
            if user:
                client.force_authenticate(user)
            res = client.post("/api/users/import", {"rows": [], "dryRun": False}, format="json")
            self.assertEqual(res.status_code, expected)

    def test_the_welcome_link_sets_a_password_and_verifies_the_email(self):
        self.post([{"name": "Nadia Salem", "email": "nadia@uni.dz", "student_id": "S100"}], dryRun=False)
        body = mail.outbox[0].body
        uid, token = re.search(r"/reset-password\?uid=([\w-]+)&token=([\w-]+)", body).groups()
        anonymous = APIClient()
        res = anonymous.post("/api/auth/password-reset/confirm", {"uid": uid, "token": token, "password": PAST},
                             format="json")
        self.assertEqual(res.status_code, 204)
        user = User.objects.get(email="nadia@uni.dz")
        self.assertTrue(user.check_password(PAST))
        self.assertTrue(user.profile.email_verified)
        self.assertEqual(anonymous.post("/api/auth/login", {"username": "nadia@uni.dz", "password": PAST},
                                        format="json").status_code, 200)
        # single use
        res = anonymous.post("/api/auth/password-reset/confirm", {"uid": uid, "token": token, "password": PAST + "x"},
                             format="json")
        self.assertEqual(res.status_code, 400)

    def test_a_reset_email_also_verifies_an_unverified_self_registered_account(self):
        user = make_user("self@x.com")
        UserProfile.objects.filter(user=user).update(email_verified=False)
        mail.outbox.clear()
        anonymous = APIClient()
        anonymous.post("/api/auth/password-reset", {"email": "self@x.com"}, format="json")
        uid, token = re.search(r"uid=([\w-]+)&token=([\w-]+)", mail.outbox[0].body).groups()
        anonymous.post("/api/auth/password-reset/confirm", {"uid": uid, "token": token, "password": PAST}, format="json")
        user.profile.refresh_from_db()
        self.assertTrue(user.profile.email_verified)

    @override_settings(EMAIL_BACKEND="library.test_exports_and_students.BrokenBackend")
    def test_a_mail_outage_does_not_undo_the_accounts(self):
        res = self.post([{"name": "Nadia Salem", "email": "nadia@uni.dz", "student_id": "S100"}], dryRun=False)
        self.assertEqual((res.status_code, res.data["summary"]["create"], res.data["emailsSent"]), (200, 1, 0))
        self.assertTrue(User.objects.filter(email="nadia@uni.dz").exists())


from django.core.mail.backends.base import BaseEmailBackend  # noqa: E402


class BrokenBackend(BaseEmailBackend):
    def send_messages(self, email_messages):
        raise OSError("SMTP is down")


class StudentImportCommandTests(StudentImportBase):
    def test_dry_run_then_import(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp, "students.csv")
            path.write_text("nom;prenom;email;matricule\nSalem;Nadia;nadia@uni.dz;S100\nBad;Row;nope;S101\n",
                            encoding="utf-8-sig")
            out = io.StringIO()
            call_command("import_students", str(path), "--dry-run", stdout=out)
            self.assertIn("Would create: 1 accounts", out.getvalue())
            self.assertIn("Row 3: Email is not valid.", out.getvalue())
            self.assertFalse(User.objects.filter(email="nadia@uni.dz").exists())
            self.assertEqual(len(mail.outbox), 0)

            out = io.StringIO()
            call_command("import_students", str(path), stdout=out)
            self.assertIn("Created: 1 accounts", out.getvalue())
            self.assertTrue(User.objects.filter(email="nadia@uni.dz").exists())
            self.assertEqual(len(mail.outbox), 1)

    def test_reads_json_and_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            wb.active.append(["Name", "Email", "Student ID"])
            wb.active.append(["Omar Ben Ali", "omar@uni.dz", 20240001])  # the ID is stored as a number
            wb.save(Path(tmp, "s.xlsx"))
            Path(tmp, "s.json").write_text('{"students": [{"name": "Lina H", "email": "lina@uni.dz", "student_id": "J1"}]}')
            for name in ("s.xlsx", "s.json"):
                call_command("import_students", str(Path(tmp, name)), stdout=io.StringIO())
        self.assertEqual(UserProfile.objects.get(user__email="omar@uni.dz").student_or_staff_id, "20240001")
        self.assertTrue(User.objects.filter(email="lina@uni.dz").exists())


class HeaderAliasTests(TestCase):
    def test_multi_word_arabic_headers_stay_distinct(self):
        """Accents are stripped but Arabic letters kept: spaces alone must not make every Arabic header the same key."""
        from .importing import map_columns
        columns = map_columns(["عنوان الكتاب", "عدد النسخ", "رمز الرف", "دار النشر"])
        self.assertEqual(list(columns.values()), ["title", "totalCopies", "shelfLocation", "publisher"])
