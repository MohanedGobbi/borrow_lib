"""Bulk import of books from spreadsheet-like rows (Excel, CSV or JSON).

The same rules back the staff import screen (POST /api/books/import) and the
`import_books` management command:

- Column headers are matched case- and accent-insensitively against English,
  French and Arabic aliases; unknown columns are ignored.
- The ISBN identifies a book when there is one: it is normalized (digits and a final X only),
  its checksum is verified, and it matches existing books regardless of hyphens. Without an ISBN
  (most of the faculty's collection) the call number and the title identify it, or the title alone.
- Without a category column, the category comes from the call number (catalogue.subject_of).
- Rows with errors are reported and skipped. Valid rows are written in one
  transaction, so a database failure leaves nothing half-imported.
- `dry_run` validates and reports without writing anything (the preview).
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import unicodedata
from dataclasses import dataclass, field
from datetime import date

from django.db import transaction

from .catalogue import GENERAL, SUBJECTS, subject_of
from .errors import BusinessRuleError
from .models import Book, Category, Status, log_activity

MAX_ROWS = 5000
MAX_COPIES = 1000
# Library binding cloth colors for generated covers (same set as the book editor).
BINDINGS = ["#1F4D3A", "#6E2A2A", "#2B3F52", "#4A4E45", "#5C4630", "#29463F"]

# Row errors are returned as codes (+ values) so the UI can show them in FR / AR / EN.
ERROR_TEXT = {
    "isbn_invalid": "ISBN is not valid (check the digits).",
    "isbn_duplicate": "Same ISBN as row {row}.",
    "book_duplicate": "Same call number and title as row {row}.",
    "title_missing": "Title is missing.",
    "title_too_long": "Title is longer than 500 characters.",
    "year_invalid": "Year must be a year such as 2019.",
    "copies_invalid": "Copies must be a whole number from 1 to {max}.",
    "copies_below_loans": "Copies cannot be below the {onLoan} on loan.",
}


class ImportError_(BusinessRuleError, ValueError):
    """Bad import request. Also a ValueError, so the management command can catch it with the file-reading errors."""


def error_text(error: dict) -> str:
    """English wording for an error code (used by the management command)."""
    return ERROR_TEXT[error["code"]].format(**error)


FIELDS = ["isbn", "title", "author", "category", "publisher", "year", "totalCopies", "shelfLocation", "description"]

_ALIASES = {
    "isbn": ["isbn", "isbn13", "isbn 13", "isbn10", "code isbn", "ردمك"],
    "title": ["title", "book title", "titre", "titre du livre", "العنوان", "عنوان الكتاب", "اسم الكتاب"],
    "author": ["author", "authors", "auteur", "auteurs", "writer", "المؤلف", "المؤلفون"],
    "category": ["category", "categories", "categorie", "genre", "subject", "discipline", "التصنيف", "الفئة"],
    "publisher": ["publisher", "editeur", "edition", "maison d edition", "الناشر", "دار النشر"],
    "year": ["year", "publication year", "published", "annee", "annee de publication", "السنة", "سنة النشر"],
    "totalCopies": ["total copies", "totalcopies", "copies", "copy count", "quantity", "qty", "exemplaires",
                    "nombre d exemplaires", "nb exemplaires", "عدد النسخ", "النسخ"],
    "shelfLocation": ["shelf location", "shelflocation", "shelf", "shelf mark", "location", "call number",
                      "cote", "emplacement", "rayon", "الرف", "رمز الرف", "الشفرة", "الرمز", "رقم الطلب",
                      "رمز التصنيف", "التصنيف العشري"],
    "description": ["description", "summary", "resume", "abstract", "notes", "الوصف", "الملخص"],
}


def _key(header: str) -> str:
    """Lowercase, strip accents and separators: 'Nombre d’exemplaires' → 'nombre d exemplaires'."""
    # Drop accents (and Arabic vowel marks) but keep the letters, so Arabic headers stay distinct from each other.
    text = "".join(ch for ch in unicodedata.normalize("NFKD", str(header)) if not unicodedata.combining(ch)).lower()
    text = text.replace("\u0640", "")  # the Arabic tatweel that stretches headers such as "الـعنـوان"
    for ch in "_-./’'()":
        text = text.replace(ch, " ")
    return " ".join(text.split())


_LOOKUP = {_key(alias): name for name, aliases in _ALIASES.items() for alias in aliases}


def map_columns(headers) -> dict[str, str | None]:
    """Map each raw header to a field name, or None when the column is not used."""
    return {h: _LOOKUP.get(_key(h)) for h in headers}


def normalize_isbn(raw) -> str | None:
    """Return the ISBN as bare digits (plus a final X for ISBN-10), or None if invalid."""
    if isinstance(raw, float) and raw.is_integer():  # Excel stores ISBN-13 as a number
        raw = int(raw)
    text = "".join(ch for ch in str(raw).upper() if ch.isdigit() or ch == "X")
    if len(text) == 13 and text.isdigit():
        total = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(text[:12]))
        return text if (10 - total % 10) % 10 == int(text[12]) else None
    if len(text) == 10 and text[:9].isdigit():
        total = sum(int(d) * (10 - i) for i, d in enumerate(text[:9]))
        check = (11 - total % 11) % 11
        return text if text[9] == ("X" if check == 10 else str(check)) else None
    return None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _int(value) -> int | None:
    text = _text(value)
    if not text:
        return None
    try:
        number = float(text.replace(",", "."))
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


@dataclass
class RowResult:
    row: int
    status: str  # create | update | unchanged | skip | error
    isbn: str = ""
    title: str = ""
    author: str = ""
    totalCopies: int | None = None
    errors: list[dict] = field(default_factory=list)
    changes: list[str] = field(default_factory=list)


def run_import(rows: list[dict], *, mode: str = "skip", dry_run: bool = True,
               actor_name: str = "System", source: str = "") -> dict:
    """Validate (and unless dry_run, apply) an import. `rows` are dicts keyed by raw headers."""
    if mode not in ("skip", "update"):
        raise ImportError_("import_bad_mode", "mode must be 'skip' or 'update'.")
    if len(rows) > MAX_ROWS:
        raise ImportError_("import_too_many_rows", f"A single import is limited to {MAX_ROWS} rows; split the file.",
                           max=MAX_ROWS)

    headers = list(dict.fromkeys(h for r in rows for h in r.keys() if not str(h).startswith("__")))
    columns = map_columns(headers)
    missing = [f for f in ("title",) if f not in columns.values()]
    if rows and missing:
        return {"dryRun": dry_run, "columns": columns, "missingColumns": missing, "rows": [],
                "summary": _summary([], [])}

    existing = {}
    for b in Book.objects.all():
        if b.isbn:
            existing[normalize_isbn(b.isbn) or b.isbn] = b
        existing.setdefault(_shelf_key(b.shelf_location, b.title), b)
    categories = {c.name.lower(): c for c in Category.objects.all()}
    this_year = date.today().year
    seen: dict[str, int] = {}
    results: list[RowResult] = []
    plans: list[tuple[RowResult, dict, Book | None]] = []
    new_categories: list[str] = []

    for index, raw in enumerate(rows):
        data = {f: None for f in FIELDS}
        for header, value in raw.items():
            name = columns.get(header)
            if name and data[name] in (None, ""):
                data[name] = value
        row_no = raw.get("__row") or index + 1
        res = RowResult(row=row_no, status="error", title=_text(data["title"]), author=_text(data["author"]))

        # blank lines in a spreadsheet are ignored, not reported
        if not any(_text(v) for v in data.values()):
            continue

        isbn = normalize_isbn(data["isbn"]) if _text(data["isbn"]) else None
        res.isbn = isbn or _text(data["isbn"])
        shelf = _text(data["shelfLocation"])
        key = isbn or _shelf_key(shelf, res.title)
        if _text(data["isbn"]) and not isbn:
            res.errors.append({"code": "isbn_invalid"})
        elif res.title and key in seen:
            res.errors.append({"code": "isbn_duplicate" if isbn else "book_duplicate", "row": seen[key]})
        if isbn or res.title:
            seen.setdefault(key, row_no)
        if not res.title:
            res.errors.append({"code": "title_missing"})
        elif len(res.title) > 500:
            res.errors.append({"code": "title_too_long"})

        year = None
        if _text(data["year"]):
            year = _int(data["year"])
            if year is None or not 1000 <= year <= this_year + 1:
                res.errors.append({"code": "year_invalid"})
                year = None
        copies = 1
        if _text(data["totalCopies"]):
            copies = _int(data["totalCopies"])
            if copies is None or not 1 <= copies <= MAX_COPIES:
                res.errors.append({"code": "copies_invalid", "max": MAX_COPIES})
                copies = 1
        res.totalCopies = copies

        book = existing.get(key)
        if book and mode == "update":
            on_loan = book.borrow_records.filter(status__in=[Status.BORROWED, Status.OVERDUE]).count()
            if copies < on_loan:
                res.errors.append({"code": "copies_below_loans", "onLoan": on_loan})

        if res.errors:
            results.append(res)
            continue

        category_name = _text(data["category"]) or subject_of(shelf)[0]
        if category_name.lower() not in categories and category_name not in new_categories:
            new_categories.append(category_name)
        values = {
            "title": res.title, "author": res.author, "category_name": category_name,
            "publisher": _text(data["publisher"]), "year": year, "total_copies": copies,
            "shelf_location": _text(data["shelfLocation"]), "description": _text(data["description"]),
        }

        if book is None:
            res.status = "create"
        elif mode == "skip":
            res.status = "skip"
        else:
            res.changes = _changes(book, values)
            res.status = "update" if res.changes else "unchanged"
        results.append(res)
        if res.status in ("create", "update"):
            plans.append((res, values, book))

    if not dry_run and plans:
        with transaction.atomic():
            for name in new_categories:
                color = next((c for n, c in _subject_colours() if n == name), "#4A5560")
                categories[name.lower()] = Category.objects.create(name=name, color=color)
            created = updated = 0
            for res, values, book in plans:
                category = categories[values.pop("category_name").lower()]
                if book is None:
                    seed = res.isbn or f"{values['shelf_location']}|{res.title}"
                    Book.objects.create(
                        isbn=res.isbn or None, category=category, available_copies=values["total_copies"],
                        cover_color=BINDINGS[int(hashlib.md5(seed.encode()).hexdigest(), 16) % len(BINDINGS)],
                        **values,
                    )
                    created += 1
                else:
                    on_loan = book.total_copies - book.available_copies
                    for attr, value in values.items():
                        setattr(book, attr, value)
                    book.category = category
                    book.available_copies = max(0, values["total_copies"] - on_loan)
                    book.save()
                    updated += 1
            where = f" from {source}" if source else ""
            log_activity(actor_name, "BOOKS_IMPORTED",
                         f"Imported {created} new books and updated {updated}{where}.",
                         created=created, updated=updated, source=source)

    return {
        "dryRun": dry_run,
        "columns": columns,
        "missingColumns": [],
        "newCategories": new_categories,
        "rows": [r.__dict__ for r in results],
        "summary": _summary(results, new_categories),
    }


def _shelf_key(shelf: str, title: str) -> str:
    """Identity of a book without an ISBN: its call number and its title, compared loosely."""
    return "shelf:" + " ".join(f"{shelf} {title}".lower().split())


def _subject_colours():
    return [*SUBJECTS.values(), GENERAL]


def _changes(book: Book, values: dict) -> list[str]:
    current = {
        "title": book.title, "author": book.author, "category_name": book.category.name,
        "publisher": book.publisher, "year": book.year, "total_copies": book.total_copies,
        "shelf_location": book.shelf_location, "description": book.description,
    }
    labels = {"category_name": "category", "total_copies": "copies", "shelf_location": "shelf"}
    return [labels.get(k, k) for k, v in values.items()
            if (current[k].lower() if k == "category_name" else current[k]) !=
               (v.lower() if k == "category_name" else v)]


def _summary(results: list[RowResult], new_categories: list[str]) -> dict:
    counts = {s: 0 for s in ("create", "update", "unchanged", "skip", "error")}
    for r in results:
        counts[r.status] += 1
    return {"total": len(results), **counts, "newCategories": len(new_categories)}


# ------------------------------------------------------------------ file readers

def read_rows(filename: str, content: bytes) -> list[dict]:
    """Read .xlsx, .csv or .json into dicts keyed by header; `__row` is the spreadsheet row number."""
    name = filename.lower()
    if name.endswith(".json"):
        data = json.loads(content.decode("utf-8-sig"))
        if isinstance(data, dict):
            data = data.get("books", data.get("students", data.get("rows", [])))
        if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
            raise ValueError("A JSON file must contain an array of row objects (or {\"books\": [...]}).")
        return [{**r, "__row": i + 1} for i, r in enumerate(data)]
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        return _from_table(list(reader))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        return _from_table([list(r) for r in sheet.iter_rows(values_only=True)])
    raise ValueError("Unsupported file type. Use .xlsx, .csv or .json.")


def _from_table(table: list[list]) -> list[dict]:
    if not table:
        return []
    headers = [_text(h) for h in table[0]]
    rows = []
    for i, cells in enumerate(table[1:], start=2):
        row = {h: cells[j] if j < len(cells) else None for j, h in enumerate(headers) if h}
        row["__row"] = i
        rows.append(row)
    return rows
